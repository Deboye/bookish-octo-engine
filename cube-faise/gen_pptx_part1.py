from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHP

W = Inches(13.333)
H = Inches(7.5)

# Colors
NAVY = RGBColor(0, 40, 85)
GOLD = RGBColor(200, 169, 86)
WHITE = RGBColor(255, 255, 255)
LGRAY = RGBColor(248, 249, 250)
MGRAY = RGBColor(108, 117, 125)
DTXT = RGBColor(33, 37, 41)
RED = RGBColor(231, 76, 60)
BLUE = RGBColor(52, 152, 219)
ORANGE = RGBColor(243, 156, 18)
GREEN = RGBColor(39, 174, 96)
DGREEN = RGBColor(10, 35, 16)
MGREEN = RGBColor(27, 94, 32)
BGREEN = RGBColor(46, 125, 50)
AGREEN = RGBColor(76, 175, 80)
CDBG = RGBColor(18, 55, 25)
MUTED = RGBColor(144, 191, 147)

def mk():
    p = Presentation()
    p.slide_width = W
    p.slide_height = H
    return p

def sl(p):
    return p.slides.add_slide(p.slide_layouts[6])

def rc(s, x, y, w, h, f, lc=None):
    sh = s.shapes.add_shape(SHP.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = f
    if lc:
        sh.line.color.rgb = lc
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    return sh

def tx(s, t, x, y, w, h, sz=12, b=False, c=None, a=PP_ALIGN.LEFT, fn='Calibri'):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = a
    r = p.add_run()
    r.text = t
    r.font.size = Pt(sz)
    r.font.bold = b
    r.font.name = fn
    if c: r.font.color.rgb = c
    return tb

def mtx(s, lines, x, y, w, h, a=PP_ALIGN.LEFT, fn='Calibri'):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(lines):
        txt, sz, bold, clr = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = a
        if txt == '':
            r = p.add_run()
            r.text = ''
            r.font.size = Pt(4)
        else:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.name = fn
            if clr: r.font.color.rgb = clr
    return tb

def sbg(slide, color):
    bg = slide.background
    f = bg.fill
    f.solid()
    f.fore_color.rgb = color

def hdr(s, label, num):
    rc(s, Inches(0), Inches(0), W, Inches(0.6), NAVY)
    rc(s, Inches(0), Inches(0.6), Inches(0.06), Inches(6.6), GOLD)
    tx(s, label, Inches(0.5), Inches(0.12), Inches(5), Inches(0.4), 9, True, GOLD, fn='Calibri')
    tx(s, num, Inches(12), Inches(0.12), Inches(1), Inches(0.4), 9, False, WHITE, PP_ALIGN.RIGHT)
    rc(s, Inches(0), Inches(7.2), W, Inches(0.3), NAVY)
    tx(s, 'CONFIDENTIEL  |  Cabinet Carree x CUBE', Inches(0.5), Inches(7.22), Inches(12), Inches(0.25), 7, False, GOLD)

print('Part 1: helpers defined OK')
